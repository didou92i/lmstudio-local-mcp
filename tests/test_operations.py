import json
import sys
from pathlib import Path

import jsonschema
import pytest
from docx import Document
from pypdf import PdfWriter

from lmstudio_mcp.client import Client, ConnectorError
from lmstudio_mcp.config import Settings
from lmstudio_mcp.connections import Connections, validate_url
from lmstudio_mcp.diagnostics import diagnose, journal
from lmstudio_mcp.integrations import Integrations, validate_config
from lmstudio_mcp.rag import Rag, unit
from lmstudio_mcp.sdk import configuration_mismatches, model_config, schema
from lmstudio_mcp.service import Service
from lmstudio_mcp.storage import atomic_json, read_json


@pytest.fixture
def settings(tmp_path):
    root = tmp_path / 'documents'
    root.mkdir()
    return Settings(state_dir=tmp_path / 'state', rag_roots=[root], mcp_config_path=tmp_path / 'mcp.json')


@pytest.mark.parametrize('url', ['http://evil.example', 'http://127.0.0.1@evil.test',
    'http://169.254.169.254', 'http://0.0.0.0', 'http://127.0.0.1:1234/v1', 'https://a.test/?token=secret'])
def test_connection_origins_reject_unsafe_or_ambiguous_urls(url):
    with pytest.raises(ConnectorError):
        validate_url(url)


async def test_connection_profiles_keep_secrets_out_and_guard_selection(settings, monkeypatch):
    con = Connections(settings)
    monkeypatch.setenv('LM_REMOTE_TEST_TOKEN', 'secret-value')
    await con.run('save', 'test', 'https://example.test', 'LM_REMOTE_TEST_TOKEN')
    assert 'secret-value' not in con.path.read_text()
    assert con.get('test')['token'] == 'secret-value'
    async def failed(profile):
        return {'compatible': False, 'issue': 'authentication'}
    monkeypatch.setattr(con, 'test', failed)
    with pytest.raises(ConnectorError, match='unchanged'):
        await con.run('select', 'test')
    assert con.registry()['active'] == 'local'
    with pytest.raises(ConnectorError, match='Token variables'):
        await con.run('save', 'bad', 'https://example.test', 'HOME')


async def test_remote_selection_disables_local_cli_and_missing_token_is_recoverable(settings, monkeypatch):
    con = Connections(settings)
    await con.run('save', 'remote', 'http://192.168.1.10:1234', 'LM_REMOTE_TEST')
    data = con.registry()
    data['active'] = 'remote'
    atomic_json(con.path, data)
    monkeypatch.delenv('LM_REMOTE_TEST', raising=False)
    client = Client(settings)
    try:
        result = await Service(client).run('connections', action='list')
        assert result['ok'] and result['data']['active'] == 'remote'
        client.active_local = False
        with pytest.raises(ConnectorError, match='disabled'):
            await client.cli(['server', 'stop'])
        async def healthy(profile):
            return {'compatible': True}
        monkeypatch.setattr(Connections, 'test', lambda self, p: healthy(p))
        result = await Service(client).run('connections', action='select', name='local')
        assert result['ok'] and con.get()['local']
    finally:
        await client.close()


def test_mcp_config_preserves_entries_backups_and_binds_permissions(settings):
    original = {'mcpServers': {'existing': {'command': '/usr/bin/true', 'args': ['private-argument']}}, 'other': 42}
    atomic_json(settings.mcp_config_path, original)
    integrations = Integrations(settings)
    preview = integrations.configure('upsert', 'test', {'command': '/usr/bin/true'}, ['ping'])
    assert read_json(settings.mcp_config_path) == original
    result = integrations.configure('upsert', 'test', {'command': '/usr/bin/true'}, ['ping'], True, preview['expected_digest'])
    assert result['verified'] and read_json(Path(result['backup'])) == original
    assert read_json(settings.mcp_config_path)['other'] == 42
    assert integrations.authorized_tools('test') == ['ping']
    assert 'private-argument' not in json.dumps(integrations.configure('list'))
    with pytest.raises(ConnectorError, match='changed'):
        integrations.configure('remove', 'test', apply=True, expected_digest=preview['expected_digest'])
    changed = read_json(settings.mcp_config_path)
    changed['mcpServers']['test']['args'] = ['changed']
    atomic_json(settings.mcp_config_path, changed)
    assert integrations.authorized_tools('test') == []


@pytest.mark.parametrize('config', [{'command': 'x', 'url': 'http://localhost'},
    {'command': 'x', 'args': 'bad'}, {'url': 'https://a.test', 'headers': {'Authorization': 5}},
    {'url': 'https://token:secret@a.test'}, {'command': 'x', 'env': {'KEY': 42}}])
def test_mcp_invalid_configuration(config):
    with pytest.raises(ConnectorError):
        validate_config(config)


async def test_real_child_mcp_handshake_authorization_schema_and_error(settings, tmp_path):
    fixture = tmp_path / 'fixture.py'
    fixture.write_text('from mcp.server.fastmcp import FastMCP\nm=FastMCP("fixture")\n'
        '@m.tool()\ndef ping(value: int) -> dict:\n return {"value":value+1}\n'
        '@m.tool()\ndef failure() -> str:\n raise ValueError("expected fixture error")\n'
        'm.run()\n')
    integrations = Integrations(settings)
    config = {'command': sys.executable, 'args': [str(fixture)]}
    preview = integrations.configure('upsert', 'fixture', config, ['ping', 'failure'])
    integrations.configure('upsert', 'fixture', config, ['ping', 'failure'], True, preview['expected_digest'])
    listing = await integrations.inspect_or_call('fixture')
    assert listing['connected'] and {t['name'] for t in listing['tools']} == {'ping', 'failure'}
    called = await integrations.inspect_or_call('fixture', 'ping', {'value': 41})
    payload = called['result'].get('structuredContent') or json.loads(called['result']['content'][0]['text'])
    assert payload['value'] == 42
    assert (await integrations.inspect_or_call('fixture', 'failure'))['status'] == 'failed'
    with pytest.raises(ConnectorError, match='not authorized'):
        await integrations.inspect_or_call('fixture', 'ungranted')
    with pytest.raises(ConnectorError, match='ValidationError'):
        await integrations.inspect_or_call('fixture', 'ping', {'value': 'wrong'})


class RagClient:
    def __init__(self, settings):
        self.settings = settings
        self.active_url = settings.base_url
        self.current_health = {'runtime_fingerprint': 'runtime1'}
        self.rows = [{'key': 'embed', 'type': 'embedding', 'loaded_instances': [], 'format': 'gguf'},
                     {'key': 'chat', 'type': 'llm', 'loaded_instances': []}]
        self.model = Client.model
        self.fail = False
        self.dimension = 2
        self.hook = None
        self.generated = {'supported': True, 'answer': 'Le code est ZEPHYR-842 [S1].',
                          'citations': [{'source_id': 'S1', 'quote': 'Le code est ZEPHYR-842.'}]}

    async def request(self, method, endpoint, payload):
        if self.fail:
            raise ConnectorError('Embedding failure')
        if endpoint == '/v1/chat/completions':
            assert payload['response_format']['type'] == 'json_schema' and 'tools' not in payload
            return {'choices': [{'message': {'content': json.dumps(self.generated)}}]}
        if self.hook:
            self.hook()
        return {'data': [{'index': i, 'embedding': [1.0] * self.dimension} for i, _ in enumerate(payload['input'])]}


@pytest.fixture
def rag(settings):
    return Rag(RagClient(settings))


async def build(rag):
    doc = rag.client.settings.rag_roots[0] / 'note.txt'
    doc.write_text('Le code est ZEPHYR-842. Le centre ferme à 18 heures.')
    result = await rag.index(rag.client.rows, 'test', [str(doc)], 'embed')
    assert result['indexed_chunks'] == 1
    return doc


async def test_rag_source_provenance_and_stale_document_exclusion(rag):
    doc = await build(rag)
    found = await rag.search(rag.client.rows, 'test', 'code?', 'embed')
    assert found['sources'][0]['path'] == str(doc) and found['sources'][0]['sha256']
    result = await rag.ask(rag.client.rows, 'test', 'code?', 'embed', 'chat')
    assert result['supported'] and '[S1]' in result['answer']
    doc.write_text('changed')
    result = await rag.ask(rag.client.rows, 'test', 'code?', 'embed', 'chat')
    assert result['answer'] is None and result['excluded_stale_documents'] == [str(doc)]
    rag.manage('delete', 'test')
    assert doc.exists() and not rag.manage('list')['collections']


async def test_rag_renders_missing_markers_only_after_validating_quotes(rag):
    await build(rag)
    rag.client.generated['answer'] = 'Le code est ZEPHYR-842.'
    result = await rag.ask(rag.client.rows, 'test', 'code?', 'embed', 'chat')
    assert result['supported'] and result['answer'].endswith('[S1]') and result['citation_markers_added']
    rag.client.generated['citations'][0]['quote'] = 'invented reference'
    result = await rag.ask(rag.client.rows, 'test', 'code?', 'embed', 'chat')
    assert result['answer'] is None


@pytest.mark.parametrize('generation', [
    {'supported': True, 'answer': 'Invented [S1]', 'citations': [{'source_id': 'S1', 'quote': 'not in document'}]},
    {'supported': True, 'answer': 'Le code est ZEPHYR-842 [S9].', 'citations': [{'source_id': 'S1', 'quote': 'Le code est ZEPHYR-842.'}]},
    {'supported': False, 'answer': 'invented'}, ['malformed'], {'supported': True, 'answer': 42},
])
async def test_rag_refuses_unverifiable_answers(rag, generation):
    await build(rag)
    rag.client.generated = generation
    result = await rag.ask(rag.client.rows, 'test', 'code?', 'embed', 'chat')
    assert result['answer'] is None and not result['supported']


async def test_rag_failed_embedding_and_concurrent_change_preserve_index(rag):
    doc = await build(rag)
    before = rag.manage('list')
    rag.client.fail = True
    with pytest.raises(ConnectorError, match='failure'):
        await rag.index(rag.client.rows, 'test', [str(doc)], 'embed', True)
    assert rag.manage('list') == before
    rag.client.fail = False
    def change():
        conn = rag.db()
        with conn:
            conn.execute("UPDATE collections SET updated_at='concurrent' WHERE name='test'")
        conn.close()
    rag.client.hook = change
    with pytest.raises(ConnectorError, match='changed during'):
        await rag.index(rag.client.rows, 'test', [str(doc)], 'embed')
    assert rag.manage('list')['collections'][0]['chunks'] == 1


async def test_rag_rejects_changed_embedding_space(rag):
    doc = await build(rag)
    rag.client.dimension = 3
    with pytest.raises(ConnectorError, match='dimension'):
        await rag.index(rag.client.rows, 'test', [str(doc)], 'embed')
    with pytest.raises(ConnectorError, match='dimension'):
        await rag.search(rag.client.rows, 'test', 'code?', 'embed')
    rag.client.current_health['runtime_fingerprint'] = 'runtime2'
    with pytest.raises(ConnectorError, match='reindex'):
        await rag.search(rag.client.rows, 'test', 'code?', 'embed')


def test_rag_document_roots_and_extractors(rag, tmp_path):
    outside = tmp_path / 'private.txt'
    outside.write_text('private')
    link = rag.client.settings.rag_roots[0] / 'link.txt'
    link.symlink_to(outside)
    for path in [outside, link]:
        with pytest.raises(ConnectorError, match='outside'):
            rag.file(path)
    docpath = rag.client.settings.rag_roots[0] / 'example.docx'
    doc = Document()
    doc.add_paragraph('Paragraph text')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Table text'
    doc.save(docpath)
    assert 'Table text' in rag.extract(docpath)[0][1]
    pdfpath = rag.client.settings.rag_roots[0] / 'empty.pdf'
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(pdfpath)
    with pytest.raises(ConnectorError, match='OCR'):
        rag.extract(pdfpath)
    pdfpath.write_bytes(b'broken')
    with pytest.raises(ConnectorError, match='cannot be parsed'):
        rag.extract(pdfpath)


@pytest.mark.parametrize('vector', [[0, 0], [float('nan')], [float('inf')], ['bad'], []])
def test_invalid_vectors(vector):
    with pytest.raises(ConnectorError):
        unit(vector)


async def test_sdk_schema_rejects_ignored_fields_and_inspect_never_loads(settings):
    jsonschema.validate({'contextLength': 4096, 'flashAttention': True}, schema())
    with pytest.raises(jsonschema.ValidationError):
        jsonschema.validate({'pretendSetting': True}, schema())
    c = Client(settings)
    try:
        with pytest.raises(ConnectorError, match='no implicit loading'):
            await model_config(c, [], 'inspect', instance_id='absent')
    finally:
        await c.close()


def test_sdk_nested_settings_compare_requested_fields_only():
    assert not configuration_mismatches({'gpu': {'ratio': 0.5}}, {'gpu': {'ratio': 0.5, 'mainGpu': 0}})
    assert configuration_mismatches({'gpu': {'ratio': 0.5}}, {'gpu': {'ratio': 1}}) == {
        'gpu.ratio': {'requested': 0.5, 'actual': 1}}


async def test_diagnostics_returns_signals_without_raw_logs(settings, tmp_path, monkeypatch):
    root = tmp_path / '.lmstudio/server-logs/2026-09'
    root.mkdir(parents=True)
    (root / 'test.log').write_text('[ERROR] out of memory PRIVATE-PROMPT\n[DEBUG] another-secret')
    monkeypatch.setattr(Path, 'home', lambda: tmp_path)
    client = Client(settings)
    async def fail(*args, **kwargs):
        raise ConnectorError('Server unreachable')
    monkeypatch.setattr(client, 'cli', fail)
    journal(client, 'chat', 'Context configuration mismatch')
    try:
        result = await diagnose(client, [])
        assert result['server']['error'] == 'Server unreachable'
        assert result['log_signals'][0]['category_counts']['memory'] == 1
        assert 'PRIVATE-PROMPT' not in json.dumps(result) and 'another-secret' not in json.dumps(result)
    finally:
        await client.close()
